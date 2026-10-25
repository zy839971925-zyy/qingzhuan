import os,pathlib,subprocess,shutil
R=pathlib.Path(__file__).resolve().parents[1];os.chdir(R)
J=os.environ.get('JAVA_HOME','/usr/lib/jvm/java-17-openjdk-amd64');os.environ['LD_LIBRARY_PATH']=J+'/lib:'+J+'/lib/server'
(R/'build/testclasses').mkdir(parents=True,exist_ok=True)
cp=':'.join([str(R/'tests/json-jvm.jar'),str(R/'build/testclasses'),str(R/'build/classes'),os.environ.get('ANDROID_JAR',str(pathlib.Path(os.environ.get('ANDROID_SDK_ROOT',str(R.parent/'tooling/sdk')))/'platforms/android-36/android.jar'))]+list(map(str,(R/'libs').glob('*.jar'))))
shutil.rmtree(R/'build/test-data',ignore_errors=True)
subprocess.run([J+'/bin/java','-m','jdk.compiler/com.sun.tools.javac.Main','-encoding','UTF-8','-cp',cp,'-d','build/testclasses','tests/EngineTest.java']+list(map(str,(R/'tests/stubs').rglob('*.java'))),check=True)
subprocess.run([J+'/bin/java','-cp',cp,'cn.qingzhuan.app.EngineTest'],check=True)
# Independent implementation reopens produced Office files.
from openpyxl import load_workbook
from lxml import etree
import zipfile,json,hashlib
for x in (R/'build/test-data').glob('sample*.xlsx'):
 w=load_workbook(x);assert w.active.max_row>=1;w.close()
with zipfile.ZipFile(R/'build/test-data/document.docx') as z:
 d=etree.fromstring(z.read('word/document.xml'));assert d.xpath('count(//*[local-name()="p"])')==3
result=json.loads((R/'build/test-data/result.json').read_text())
assert result['status']=='passed'
for d in [0,1,25,26,27,701,702,16383]:
 n=d+1;s=''
 while n:n,r=divmod(n-1,26);s=chr(65+r)+s
 back=0
 for c in s:back=back*26+ord(c)-64
 assert back-1==d
print('Python independent Office validation and boundaries passed')

from pypdf import PdfReader
pdf=PdfReader(R/'build/test-data/pdf/合并文档.pdf');assert len(pdf.pages)==4
for ix,page in enumerate(pdf.pages):
 assert len(page.images)==1
 assert f'Document {ix//2+1} page {ix%2+1}' in page.extract_text()
split=PdfReader(R/'build/test-data/splitpdf/doc1-p2.pdf');assert len(split.pages)==1 and 'page 2' in split.pages[0].extract_text()
print('Python independent PDF page, text, image resources validation passed')
